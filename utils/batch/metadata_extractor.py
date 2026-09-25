"""
Metadata Extractor for extracting experimental metadata from file formats.

This module extracts metadata from various file formats (.czi, .tiff_metadata.xml,
.fcs, .xlsx, .wsp, .xml) without losing any information.

Domain-specific knowledge (FCS marker mappings, known operators, instrument
aliases) is loaded from the active profile via
:func:`utils.config_loader.get_profile`, making the extractor reusable
for any project.
"""

import hashlib
import json
import logging
import os
import re
import xml.etree.ElementTree as ET
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from utils.config_loader import get_profile

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Excel (XLSX/XLSM) structure helpers — shared single source of truth for
# :meth:`MetadataExtractor.extract_excel_metadata` and
# :meth:`FormatConverter.convert_excel_to_csv`
# (see plans/xlsx-conversion-plan.md)
# ---------------------------------------------------------------------------

try:
    import openpyxl
except ImportError:  # pragma: no cover - openpyxl is a core dependency
    openpyxl = None  # type: ignore[assignment, misc]

# Concentration units recognised in ``number + unit`` cells.  Keys are the
# lowercase spellings accepted in cells/headers (``µ`` and ``U`` are both
# accepted); values are the canonical unit reported in the metadata sidecar.
_CONC_UNIT_ALIASES = {
    "ng/ul": "ng/ul",
    "ng/ml": "ng/ml",
    "mg/ml": "mg/ml",
    "ug/ml": "ug/ml",
    "µg/ml": "ug/ml",
    "um": "um",
    "µm": "um",
    "nm": "nm",
    "mm": "mm",
    "ug": "ug",
    "µg": "ug",
    "ng": "ng",
    "mg": "mg",
    "%": "%",
}

# Longest first so e.g. "mg/ml" wins over "mg" in alternation.
_CONC_UNITS = sorted({key.lower() for key in _CONC_UNIT_ALIASES}, key=len, reverse=True)

_CONC_VALUE_RE = re.compile(
    r"^\s*([-+]?(?:\d+\.?\d*|\.\d+))\s*(" + "|".join(_CONC_UNITS) + r")\s*$",
    re.IGNORECASE,
)

_CONC_UNIT_TAIL_RE = re.compile(
    r"(" + "|".join(_CONC_UNITS) + r")\s*[\)\]\}.:]*\s*$", re.IGNORECASE
)


def _clean_cell(value: Any) -> str:
    """Normalise one worksheet cell to its string form ("" for ``None``).

    Ints are written as-is; floats via Python's shortest round-trip
    representation (``0.5`` → ``"0.5"``); integral floats are written as
    ints (``42.0`` → ``"42"``).
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


def _split_concentration(value: Any) -> Optional[Tuple[float, str]]:
    """Split a ``number + unit`` cell value into ``(number, canonical_unit)``.

    Returns ``None`` when the value is not a number followed by a known
    concentration unit.
    """
    if value is None or isinstance(value, (int, float, bool)):
        return None
    match = _CONC_VALUE_RE.match(str(value))
    if not match:
        return None
    try:
        number = float(match.group(1))
    except ValueError:
        return None
    return number, _CONC_UNIT_ALIASES[match.group(2).lower()]


def _header_unit(header: str) -> str:
    """Extract a concentration unit from the tail of a header, if any."""
    match = _CONC_UNIT_TAIL_RE.search(header.lower())
    if not match:
        return ""
    return _CONC_UNIT_ALIASES[match.group(1).lower()]


def _count_numeric(rows: List[Tuple[Any, ...]]) -> int:
    """Count numeric cells across ``rows`` (bools excluded)."""
    return sum(
        1
        for row in rows
        for cell in row
        if isinstance(cell, (int, float)) and not isinstance(cell, bool)
    )


def _detect_concentrations(
    headers: List[str], data_rows: List[Tuple[Any, ...]]
) -> List[Dict[str, Any]]:
    """Best-effort concentration detection for one sheet (plan §3.1).

    A column qualifies when its header contains ``conc`` (case-insensitive)
    or when its cells look like ``number + unit`` with a known concentration
    unit.  One entry ``{header, value, unit, row}`` is recorded per
    qualifying cell (``row`` is the 1-based data-row index).
    """
    found: List[Dict[str, Any]] = []
    if not headers or not data_rows:
        return found
    for col_idx, header in enumerate(headers):
        if not header:
            continue
        header_match = "conc" in header.lower()
        fallback_unit = _header_unit(header)
        for row_idx, row in enumerate(data_rows, start=1):
            if col_idx >= len(row):
                continue
            value = row[col_idx]
            split = _split_concentration(value)
            if split is not None:
                number, unit = split
            elif header_match and isinstance(value, (int, float)) and not isinstance(value, bool):
                number, unit = float(value), fallback_unit or "unknown"
            else:
                continue
            found.append({"header": header, "value": number, "unit": unit, "row": row_idx})
    return found


def _infer_sheet_role(name: str, headers: List[str], data_rows: List[Tuple[Any, ...]]) -> str:
    """Heuristic sheet role: ``data`` / ``calibration`` / ``metadata`` / ``unknown``."""
    if not headers and not data_rows:
        return "unknown"
    text = " ".join([name] + [h for h in headers if h]).lower()
    if re.search(r"calibrat|standard|y\s*=\s*m\s*x\s*\+\s*b|curve", text):
        return "calibration"
    if re.search(r"\b(metadata|notes?|info)\b", text):
        return "metadata"
    if _count_numeric(data_rows) >= 3:
        return "data"
    if any("conc" in h.lower() or _split_concentration(h) for h in headers):
        return "data"
    if data_rows:
        return "metadata"
    return "unknown"


def _score_primary_sheet(name: str, headers: List[str], data_rows: List[Tuple[Any, ...]]) -> int:
    """Score a sheet for primary-sheet selection (plan §3.4, deterministic).

    - ``+3`` header row has ≥ 2 non-empty string cells (looks like a table)
    - ``+2`` has ≥ 3 numeric data values
    - ``+2`` name matches ``data|result|readings?|measurement|sample``
    (the ``+1`` most-data-rows bonus is applied cross-sheet by
    :func:`read_excel_structure`).
    """
    score = 0
    if sum(1 for h in headers if _clean_cell(h)) >= 2:
        score += 3
    if _count_numeric(data_rows) >= 3:
        score += 2
    if re.search(r"data|result|readings?|measurement|sample", name, re.IGNORECASE):
        score += 2
    return score


def _count_uncached_formulas(
    file_path: str,
    sheet_name: str,
    header_excel_row: Optional[int],
    data_rows: List[Tuple[Any, ...]],
) -> int:
    """Best-effort count of formula cells whose cached value is missing.

    With ``data_only=True`` a formula cell yields its cached computed value;
    when the workbook was never opened in Excel the cache is absent and the
    cell reads as ``None`` (never the formula string).  A second read-only
    pass in formula mode identifies the formula cells, which are then
    cross-checked against the cached values already read.  Returns 0 on any
    problem (best-effort, never raises).
    """
    if openpyxl is None or header_excel_row is None or not data_rows:
        return 0
    try:
        workbook = openpyxl.load_workbook(str(file_path), read_only=True)
    except Exception:  # noqa: BLE001 - best-effort helper
        return 0
    try:
        worksheet = workbook[sheet_name]
        count = 0
        for excel_row_idx, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if excel_row_idx <= header_excel_row:
                continue
            data_row_offset = excel_row_idx - header_excel_row - 1
            if data_row_offset >= len(data_rows):
                break
            cached_row = data_rows[data_row_offset]
            for col_idx, value in enumerate(row, start=1):
                if not (isinstance(value, str) and value.lstrip().startswith("=")):
                    continue
                if col_idx - 1 >= len(cached_row) or cached_row[col_idx - 1] is None:
                    count += 1
        return count
    except Exception:  # noqa: BLE001 - best-effort helper
        return 0
    finally:
        workbook.close()


def _detect_excel_format(file_path: str) -> str:
    """Detect the Excel container format from the file's magic bytes.

    Reads exactly 4 bytes (the historical bug compared an 8-byte read
    against a 4-byte signature, so real ``.xlsx`` files reported
    "Unknown"); a ``.xlsm`` extension is reported as XLSM.
    """
    if os.path.splitext(str(file_path))[1].lower() == ".xlsm":
        return "XLSM (ZIP-based)"
    try:
        with open(file_path, "rb") as f:
            signature = f.read(4)
    except OSError:
        return "Unknown"
    if signature[:4] == b"\x50\x4b\x03\x04":
        return "XLSX (ZIP-based)"
    if signature[:2] == b"\xd0\xcf":
        return "XLS (OLE2)"
    return "Unknown"


def read_excel_structure(file_path: str) -> Dict[str, Any]:
    """Open an ``.xlsx``/``.xlsm`` workbook read-only and return its structure.

    Shared single source of truth for :meth:`MetadataExtractor.extract_excel_metadata`
    and :meth:`FormatConverter.convert_excel_to_csv`.  Uses
    ``openpyxl.load_workbook(read_only=True, data_only=True)`` so formula
    cells yield their cached computed value (never the formula string) and
    large workbooks stream with low memory.

    Never raises: every failure (missing file, zero-byte file, corrupt or
    truncated ZIP, password-protected workbook, missing openpyxl, ...) is
    reported in the ``error`` field so callers can degrade gracefully.

    Args:
        file_path: Path to the Excel file.

    Returns:
        ``{"ok": bool, "error": Optional[str], "format": str,
        "sheet_names": List[str], "sheets": List[dict],
        "primary_index": int, "concentrations": List[dict]}`` where each
        sheet entry has ``name``, ``headers``, ``rows`` (raw data rows),
        ``row_count``, ``column_count``, ``role``,
        ``detected_concentrations`` and ``data_range``.
    """
    result: Dict[str, Any] = {
        "ok": False,
        "error": None,
        "format": _detect_excel_format(file_path),
        "sheet_names": [],
        "sheets": [],
        "primary_index": 0,
        "concentrations": [],
    }

    try:
        if not os.path.exists(file_path):
            result["error"] = f"file not found: {file_path}"
            return result
        if os.stat(file_path).st_size == 0:
            result["error"] = "zero-byte file"
            return result
    except OSError as e:
        result["error"] = f"cannot access file: {e}"
        return result

    if openpyxl is None:
        result["error"] = "openpyxl is required – install with: pip install openpyxl"
        return result

    try:
        workbook = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - openpyxl raises many exception types
        if result["format"] == "XLS (OLE2)":
            result["error"] = "legacy .xls (OLE2) is not supported by openpyxl – export to .xlsx"
        else:
            result["error"] = (
                f"openpyxl could not open the file ({e}); "
                "the workbook may be corrupt, truncated, or password-protected"
            )
        return result

    sheets: List[Dict[str, Any]] = []
    sheet_names: List[str] = []
    try:
        for worksheet in workbook.worksheets:
            name = worksheet.title
            # Keep only non-empty rows together with their 1-based Excel row
            # number (needed for data_range); merged non-top-left cells
            # surface as None here and are written as empty CSV cells.
            rows_with_idx: List[Tuple[int, Tuple[Any, ...]]] = []
            for excel_row_idx, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if row and any(_clean_cell(cell) != "" for cell in row):
                    rows_with_idx.append((excel_row_idx, row))

            headers: List[str] = []
            data_rows: List[Tuple[Any, ...]] = []
            first_row: Optional[int] = None
            last_row: Optional[int] = None
            uncached_formulas = 0
            if rows_with_idx:
                first_row = rows_with_idx[0][0]
                last_row = rows_with_idx[-1][0]
                headers = [_clean_cell(cell) for cell in rows_with_idx[0][1]]
                data_rows = [row for _, row in rows_with_idx[1:]]
                uncached_formulas = _count_uncached_formulas(file_path, name, first_row, data_rows)

            column_count = max((len(row) for _, row in rows_with_idx), default=0)
            detected = _detect_concentrations(headers, data_rows)
            sheet_names.append(name)
            sheets.append(
                {
                    "name": name,
                    "headers": headers,
                    "rows": data_rows,
                    "row_count": len(data_rows),
                    "column_count": column_count,
                    "role": _infer_sheet_role(name, headers, data_rows),
                    "detected_concentrations": detected,
                    "uncached_formula_cells": uncached_formulas,
                    "data_range": {
                        "first_row": first_row,
                        "last_row": last_row,
                        "first_col": 1,
                        "last_col": column_count,
                    },
                }
            )
    except Exception as e:  # noqa: BLE001 - keep the never-raises contract
        result["error"] = f"failed while reading workbook: {e}"
        return result
    finally:
        workbook.close()

    # Primary-sheet selection (plan §3.4): highest score wins, ties go to
    # the earliest sheet; the "+1 most data rows" bonus is cross-sheet.
    max_row_count = max((sheet["row_count"] for sheet in sheets), default=0)
    best_index = 0
    best_score = -1
    for idx, sheet in enumerate(sheets):
        score = _score_primary_sheet(sheet["name"], sheet["headers"], sheet["rows"])
        if max_row_count > 0 and sheet["row_count"] == max_row_count:
            score += 1
        if score > best_score:
            best_score = score
            best_index = idx

    result.update(
        {
            "ok": True,
            "sheet_names": sheet_names,
            "sheets": sheets,
            "primary_index": best_index,
            "concentrations": [
                conc for sheet in sheets for conc in sheet["detected_concentrations"]
            ],
        }
    )
    return result


@dataclass
class FileMetadata:
    """Metadata extracted from a file."""

    file_path: str
    file_name: str
    file_type: str
    file_size: int
    created_date: Optional[str] = None
    modified_date: Optional[str] = None
    raw_metadata: Dict[str, Any] = field(default_factory=dict)
    extracted_data: Dict[str, Any] = field(default_factory=dict)
    checksum: Optional[str] = None


@dataclass
class ExperimentMetadata:
    """Complete metadata for an experiment folder."""

    experiment_id: str
    experiment_name: str
    folder_path: str
    files: List[FileMetadata] = field(default_factory=list)
    parameters: Dict[str, Any] = field(default_factory=dict)
    measurements: Dict[str, Any] = field(default_factory=dict)
    equipment: Dict[str, str] = field(default_factory=dict)
    raw_metadata: Dict[str, Any] = field(default_factory=dict)


class MetadataExtractor:
    """Extractor for experimental metadata from files."""

    def __init__(self):
        """Initialize the metadata extractor."""
        self.logger = logging.getLogger(__name__)

    def calculate_checksum(self, file_path: str) -> str:
        """
        Calculate MD5 checksum of a file.

        Args:
            file_path: Path to the file

        Returns:
            MD5 checksum as hex string
        """
        md5_hash = hashlib.md5()
        with open(file_path, "rb") as f:
            for chunk in iter(lambda: f.read(4096), b""):
                md5_hash.update(chunk)
        return md5_hash.hexdigest()

    def extract_zeiss_xml_metadata(self, xml_path: str) -> Dict[str, Any]:
        """Extract instrument metadata from Zeiss *_metadata.xml files.

        Parses the XML companion files that Zeiss microscopes produce
        alongside .czi and .tif images.  Extracts objective, magnification,
        numerical aperture, exposure times, fluorescence dye information,
        acquisition timestamps, and Z-stack parameters.

        Args:
            xml_path: Path to the Zeiss XML metadata file.

        Returns:
            Dictionary of extracted metadata.  Empty dict on failure.
        """
        import xml.etree.ElementTree as ET

        metadata: Dict[str, Any] = {}
        try:
            tree = ET.parse(xml_path)
            root = tree.getroot()

            # --- Instrument / optics ---
            for elem in root.iter("ObjectiveName"):
                if elem.text and elem.text.strip():
                    metadata["objective"] = elem.text.strip()
            for elem in root.iter("Magnification"):
                if elem.text and elem.text.strip():
                    metadata["magnification"] = elem.text.strip()
            for elem in root.iter("NumericalAperture"):
                if elem.text and elem.text.strip():
                    metadata["numerical_aperture"] = elem.text.strip()
            for elem in root.iter("Immersion"):
                if elem.text and elem.text.strip():
                    metadata["immersion"] = elem.text.strip()
            for elem in root.iter("MicroscopeModel"):
                if elem.text and elem.text.strip():
                    metadata["microscope_model"] = elem.text.strip()

            # --- Z-stack ---
            for elem in root.iter("ZStackMode"):
                if elem.text and elem.text.strip():
                    metadata["z_stack_mode"] = elem.text.strip()
            for elem in root.iter("Sections"):
                if elem.text and elem.text.strip():
                    metadata["z_stack_sections"] = elem.text.strip()

            # --- Fluorescence dyes ---
            dyes: List[Dict[str, str]] = []
            for dye_elem in root.iter("FluorescenceDye"):
                dye_info: Dict[str, str] = {}
                for child in dye_elem:
                    tag = child.tag
                    if child.text and child.text.strip():
                        dye_info[tag] = child.text.strip()
                if dye_info.get("Name"):
                    dyes.append(dye_info)
            if dyes:
                metadata["dyes"] = dyes

            # --- Exposure times ---
            exposures: List[str] = []
            for elem in root.iter("ExposureTime"):
                if elem.text and elem.text.strip():
                    exposures.append(elem.text.strip())
            if exposures:
                metadata["exposure_times_ms"] = exposures

            # --- Acquisition timestamps ---
            for elem in root.iter("StartTime"):
                if elem.text and elem.text.strip():
                    metadata["acquisition_start"] = elem.text.strip()
            for elem in root.iter("EndTime"):
                if elem.text and elem.text.strip():
                    metadata["acquisition_end"] = elem.text.strip()

            self.logger.debug(f"Extracted Zeiss XML metadata from {xml_path}")
        except Exception as e:
            self.logger.warning(f"Could not parse Zeiss XML {xml_path}: {e}")

        return metadata

    def extract_czi_metadata(self, file_path: str) -> Dict[str, Any]:
        """
        Extract metadata from Zeiss CZI microscopy files.

        Args:
            file_path: Path to the CZI file

        Returns:
            Dictionary of extracted metadata
        """
        metadata = {
            "file_type": "czi",
            "extraction_method": "basic_file_info",
            "notes": "Full CZI metadata extraction requires specialized library (czifile)",
        }

        try:
            # Basic file info
            stat_info = os.stat(file_path)
            metadata["file_size"] = str(stat_info.st_size)
            metadata["created_date"] = str(stat_info.st_ctime)
            metadata["modified_date"] = str(stat_info.st_mtime)

            # Try to extract basic metadata using binary reading
            with open(file_path, "rb") as f:
                # Read first 1KB for basic info
                header = f.read(1024)
                metadata["header_preview"] = header.hex()[:200]  # First 100 bytes as hex

                # Look for common patterns
                if b"ZISRAW" in header:
                    metadata["format"] = "ZISRAW"
                if b"Zeiss" in header:
                    metadata["vendor"] = "Zeiss"

            self.logger.info(f"Extracted basic CZI metadata from {file_path}")

        except Exception as e:
            self.logger.error(f"Error extracting CZI metadata from {file_path}: {e}")
            metadata["error"] = str(e)

        return metadata

    def extract_tiff_metadata(self, file_path: str) -> Dict[str, Any]:
        """
        Extract metadata from TIFF files.

        Args:
            file_path: Path to the TIFF file

        Returns:
            Dictionary of extracted metadata
        """
        metadata = {
            "file_type": "tiff",
            "extraction_method": "basic_file_info",
            "notes": "Full TIFF metadata extraction requires specialized library (tifffile)",
        }

        try:
            # Basic file info
            stat_info = os.stat(file_path)
            metadata["file_size"] = str(stat_info.st_size)
            metadata["created_date"] = str(stat_info.st_ctime)
            metadata["modified_date"] = str(stat_info.st_mtime)

            # Try to extract basic TIFF header info
            with open(file_path, "rb") as f:
                # Read TIFF header
                header = f.read(8)
                if header[:2] == b"II":
                    metadata["byte_order"] = "little-endian"
                elif header[:2] == b"MM":
                    metadata["byte_order"] = "big-endian"
                else:
                    metadata["byte_order"] = "unknown"

                magic = int.from_bytes(header[2:4], byteorder="big")
                if magic == 42:
                    metadata["format"] = "TIFF"
                else:
                    metadata["format"] = f"Unknown (magic: {magic})"

            self.logger.info(f"Extracted basic TIFF metadata from {file_path}")

        except Exception as e:
            self.logger.error(f"Error extracting TIFF metadata from {file_path}: {e}")
            metadata["error"] = str(e)

        return metadata

    def extract_tiff_metadata_xml(self, file_path: str) -> Dict[str, Any]:
        """
        Extract metadata from TIFF metadata XML files.

        Args:
            file_path: Path to the XML metadata file

        Returns:
            Dictionary of extracted metadata
        """
        metadata = {"file_type": "tiff_metadata_xml", "extraction_method": "xml_parsing"}

        try:
            tree = ET.parse(file_path)
            root = tree.getroot()

            # Extract all elements recursively
            def extract_elements(element, prefix=""):
                result = {}
                for child in element:
                    key = f"{prefix}/{child.tag}" if prefix else child.tag
                    if len(child) > 0:
                        result[key] = extract_elements(child, key)
                    else:
                        result[key] = child.text
                return result

            metadata["xml_data"] = json.dumps(extract_elements(root), indent=2, default=str)

            # Extract specific common fields
            if "ImageMetadata" in root.tag:
                metadata["experiment"] = root.findtext(".//ExperimentBlockIndex") or ""
                metadata["image_name"] = root.findtext(".//ImageName") or ""
                auto_save_name = root.findtext(".//AutoSave/Name") or ""
                auto_save_format = root.findtext(".//AutoSave/SingleFileSaveFormat") or ""
                metadata["auto_save_name"] = auto_save_name
                metadata["auto_save_format"] = auto_save_format

            self.logger.info(f"Extracted TIFF XML metadata from {file_path}")

        except Exception as e:
            self.logger.error(f"Error extracting TIFF XML metadata from {file_path}: {e}")
            metadata["error"] = str(e)

        return metadata

    def extract_fcs_metadata(self, file_path: str) -> Dict[str, Any]:
        """
        Extract metadata from FACS flow cytometry files.

        Args:
            file_path: Path to the FCS file

        Returns:
            Dictionary of extracted metadata
        """
        metadata: Dict[str, Any] = {
            "file_type": "fcs",
            "extraction_method": "basic_file_info",
            "notes": "Full FCS metadata extraction requires specialized library (fcsparser)",
        }

        try:
            # Basic file info
            stat_info = os.stat(file_path)
            metadata["file_size"] = str(stat_info.st_size)
            metadata["created_date"] = str(stat_info.st_ctime)
            metadata["modified_date"] = str(stat_info.st_mtime)

            # --- FCS header (ASCII, fixed 58-byte layout) ---
            # Per FCS standard:
            #   bytes  0-9 : magic + version ("FCS   3.1")
            #   bytes 10-17: TEXT segment start offset (ASCII decimal)
            #   bytes 18-25: TEXT segment end offset   (ASCII decimal)
            #   bytes 26-33: DATA segment start offset
            #   bytes 34-41: DATA segment end offset
            #   bytes 42-49: ANALYSIS segment start offset
            #   bytes 50-57: ANALYSIS segment end offset
            #
            # NOTE: the previous implementation read these offsets as big-endian
            # integers from the wrong byte ranges and scanned for values with a
            # '\x00' delimiter, which produced garbage. The FCS TEXT segment is a
            # delimiter-separated key/value list where the delimiter is the first
            # byte of the segment.
            with open(file_path, "rb") as f:
                raw = f.read()
            header = raw[:58]
            metadata["fcs_version"] = header[0:6].decode("ascii", errors="ignore").strip()

            def _header_int(start: int, end: int) -> Optional[int]:
                chunk = header[start:end].decode("ascii", errors="ignore").strip()
                try:
                    return int(chunk)
                except ValueError:
                    return None

            text_start = _header_int(10, 18)
            text_end = _header_int(18, 26)
            data_start = _header_int(26, 34)
            data_end = _header_int(34, 42)
            analysis_start = _header_int(42, 50)
            analysis_end = _header_int(50, 58)

            # Backward-compatible keys (now with correct values)
            metadata["text_start"] = str(text_start) if text_start is not None else ""
            metadata["text_end"] = str(text_end) if text_end is not None else ""
            metadata["data_start"] = str(data_start) if data_start is not None else ""
            metadata["data_end"] = str(data_end) if data_end is not None else ""
            metadata["analysis_start"] = str(analysis_start) if analysis_start is not None else ""
            metadata["analysis_end"] = str(analysis_end) if analysis_end is not None else ""

            # --- Parse the TEXT segment (delimiter-separated key/value pairs) ---
            keywords: Dict[str, str] = {}
            if (
                text_start is not None
                and text_end is not None
                and 0 < text_start < text_end <= len(raw)
            ):
                text_segment = raw[text_start:text_end].decode("latin1", errors="ignore")
                if text_segment:
                    delimiter = text_segment[0]
                    tokens = text_segment.split(delimiter)
                    # tokens -> ['', '$BEGINANALYSIS', '0', '$BEGINDATA', '256', ...]
                    for i in range(1, len(tokens) - 1, 2):
                        key = tokens[i].strip()
                        if key:
                            keywords[key] = tokens[i + 1].strip()
            metadata["keywords"] = keywords

            # Backward-compatible keyword keys (now parsed correctly)
            for keyword in ("$PAR", "$TOT", "$COMP", "$P1N", "$P1R", "$P1B"):
                if keyword in keywords:
                    metadata[keyword] = keywords[keyword]

            # Experimentally relevant fields
            # Try multiple keyword variants for instrument
            instrument = (
                keywords.get("$CYT", "")
                or keywords.get("$CYTSN", "")
                or keywords.get("$INST", "")
                or keywords.get("$CREATOR", "")
            ).strip()
            # Try multiple keyword variants for operator
            # Note: $SYS often contains the OS name (e.g. "Windows NT 6.2"),
            # not the operator.  Only use it when the value looks like a
            # person name (2-4 words, no digits, no common OS keywords).
            _op = (keywords.get("$OP", "") or keywords.get("$OPERATOR", "")).strip()
            if not _op:
                sys_val = keywords.get("$SYS", "").strip()
                # Reject OS-like values
                _os_indicators = {"windows", "linux", "macos", "darwin", "unknown"}
                if (
                    sys_val
                    and not any(kw in sys_val.lower() for kw in _os_indicators)
                    and not re.search(r"\d", sys_val)
                ):
                    _op = sys_val
            operator = _op

            # Fallback: infer operator from filename patterns using an
            # explicit allowlist to avoid false positives.
            # e.g. "HRMVEC271124JP" → operator "JP"
            known_operators = set(get_profile().get_known_operators())
            if not operator:
                fname_stem = Path(file_path).stem
                for known_op in known_operators:
                    if known_op in fname_stem:
                        operator = known_op
                        break

            # Fallback: infer instrument from context using profile aliases
            if not instrument:
                fname_lower = Path(file_path).stem.lower()
                instrument_aliases = get_profile().get_instrument_aliases()
                for alias_key, alias_value in instrument_aliases.items():
                    if alias_key in fname_lower:
                        instrument = alias_value
                        break

            metadata["instrument"] = instrument
            metadata["operator"] = operator
            metadata["acquisition_date"] = keywords.get("$DATE", "").strip()
            metadata["acquisition_date_iso"] = self._parse_fcs_date(metadata["acquisition_date"])
            metadata["begin_time"] = keywords.get("$BTIM", "").strip()
            metadata["end_time"] = keywords.get("$ETIM", "").strip()
            metadata["source"] = keywords.get("$SRC", "").strip()
            metadata["sample_number"] = keywords.get(
                "$SMNO", keywords.get("$SAMPLE_ID", "")
            ).strip()
            metadata["total_events"] = keywords.get("$TOT", "").strip()
            metadata["parameter_count"] = keywords.get("$PAR", "").strip()

            # Ordered list of $PnN channel names (e.g. FSC-A, FITC-A, PI-A)
            chan_pairs: List[Tuple[int, str]] = []
            chan_re = re.compile(r"^\$P(\d+)N$")
            for key, value in keywords.items():
                m = chan_re.match(key)
                if m:
                    chan_pairs.append((int(m.group(1)), value.strip()))
            chan_pairs.sort(key=lambda p: p[0])
            metadata["channels"] = [name for _, name in chan_pairs]
            metadata["markers"] = self._infer_fcs_markers(metadata["channels"])

            self.logger.info(f"Extracted FCS metadata from {file_path}")

        except Exception as e:
            self.logger.error(f"Error extracting FCS metadata from {file_path}: {e}")
            metadata["error"] = str(e)

        return metadata

    def _parse_fcs_date(self, date_str: str) -> str:
        """Best-effort conversion of an FCS $DATE value to ISO YYYY-MM-DD.

        Handles common FCS date formats such as "16-DEC-2025", "2025-12-16",
        "12/16/2025", and "16.12.2025". Returns an empty string when the value
        cannot be parsed.
        """
        if not date_str:
            return ""
        s = date_str.strip()
        formats = [
            "%d-%b-%Y",
            "%d-%B-%Y",
            "%d/%b/%Y",
            "%d/%B/%Y",
            "%Y-%m-%d",
            "%Y/%m/%d",
            "%d.%m.%Y",
            "%d-%m-%Y",
            "%m/%d/%Y",
        ]
        for fmt in formats:
            try:
                return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
            except ValueError:
                continue
        return ""

    def _infer_fcs_markers(self, channels: List[str]) -> List[Dict[str, str]]:
        """Infer dyes/markers from FCS channel names using profile configuration.

        Matching uses the detector token (the part before any area/height/width
        suffix such as ``-A``/``-H``/``-W``) and checks longer channel names
        first so that e.g. ``DAPI-A`` is not misread as "PI".

        The channel-to-dye mapping is loaded from the active profile's
        ``fcs_markers`` configuration.
        """
        channel_markers = get_profile().get_fcs_markers()
        # Sort by key length descending so longer names match first
        # (e.g. "DAPI" before "PI")
        sorted_keys = sorted(channel_markers, key=len, reverse=True)

        markers: List[Dict[str, str]] = []
        for ch in channels:
            detector = ch.upper().split("-")[0]
            for key in sorted_keys:
                if key.upper() in detector:
                    marker_info = channel_markers[key]
                    markers.append(
                        {
                            "channel": ch,
                            "dye": marker_info.get("dye", key),
                            "role": marker_info.get("role", ""),
                        }
                    )
                    break
        return markers

    def extract_fcs_acquisition_summary(self, folder_path: str) -> Dict[str, Any]:
        """Scan all ``.fcs`` files under *folder_path* and aggregate acquisition metadata.

        Returns a consensus dictionary with keys: ``instrument``, ``operator``,
        ``acquisition_date``, ``acquisition_date_iso``, ``channels`` (union),
        ``markers``, ``is_live_dead``, ``file_count`` and ``per_file`` details.
        Returns ``{'file_count': 0}`` when no ``.fcs`` files are found, so
        callers can gracefully fall back to filename-based behaviour.
        """
        summary: Dict[str, Any] = {"file_count": 0}
        folder = Path(folder_path)
        if not folder.exists():
            return summary

        fcs_files = sorted(
            p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() == ".fcs"
        )
        if not fcs_files:
            return summary

        per_file: List[Dict[str, Any]] = []
        instruments: List[str] = []
        operators: List[str] = []
        dates_iso: List[str] = []
        channels_union: List[str] = []
        for fp in fcs_files:
            md = self.extract_fcs_metadata(str(fp))
            inst = md.get("instrument", "")
            op = md.get("operator", "")
            date_iso = md.get("acquisition_date_iso", "")
            if inst:
                instruments.append(inst)
            if op:
                operators.append(op)
            if date_iso:
                dates_iso.append(date_iso)
            for ch in md.get("channels", []):
                if ch not in channels_union:
                    channels_union.append(ch)
            per_file.append(
                {
                    "file": fp.name,
                    "instrument": inst,
                    "operator": op,
                    "acquisition_date": md.get("acquisition_date", ""),
                    "acquisition_date_iso": date_iso,
                    "channels": md.get("channels", []),
                }
            )

        def _most_common(values: List[str]) -> str:
            nonempty = [v for v in values if v]
            return Counter(nonempty).most_common(1)[0][0] if nonempty else ""

        summary["instrument"] = _most_common(instruments)
        summary["operator"] = _most_common(operators)
        summary["acquisition_date_iso"] = _most_common(dates_iso) or (
            min(dates_iso) if dates_iso else ""
        )
        summary["acquisition_date"] = next(
            (
                pf["acquisition_date"]
                for pf in per_file
                if pf["acquisition_date_iso"] == summary["acquisition_date_iso"]
                and pf["acquisition_date"]
            ),
            "",
        )
        summary["channels"] = channels_union
        summary["markers"] = self._infer_fcs_markers(channels_union)
        summary["is_live_dead"] = any(m["dye"] == "Calcein-AM" for m in summary["markers"]) and any(
            m["dye"] == "Propidium Iodide" for m in summary["markers"]
        )
        summary["file_count"] = len(fcs_files)
        summary["per_file"] = per_file
        return summary

    def extract_excel_metadata(self, file_path: str) -> Dict[str, Any]:
        """
        Extract metadata from Excel files.

        In addition to the basic file info, the full workbook structure is
        extracted via :func:`read_excel_structure` (sheet names, per-sheet
        headers, row counts, data ranges, sheet roles and detected
        concentrations) — see plans/xlsx-conversion-plan.md.

        Args:
            file_path: Path to the Excel file

        Returns:
            Dictionary of extracted metadata
        """
        metadata: Dict[str, Any] = {
            "file_type": "excel",
            "extraction_method": "basic_file_info",
        }

        try:
            # Basic file info
            stat_info = os.stat(file_path)
            metadata["file_size"] = str(stat_info.st_size)
            metadata["created_date"] = str(stat_info.st_ctime)
            metadata["modified_date"] = str(stat_info.st_mtime)

            # Detect container format from the file signature
            # (compares exactly the first 4 bytes – the historical bug
            # compared an 8-byte read against a 4-byte signature)
            metadata["format"] = _detect_excel_format(file_path)

            # Real workbook structure (never raises; see read_excel_structure)
            structure = read_excel_structure(file_path)
            if structure["ok"]:
                metadata["extraction_method"] = "basic_file_info + openpyxl"
                metadata["sheet_count"] = len(structure["sheet_names"])
                metadata["sheet_names"] = list(structure["sheet_names"])
                metadata["sheets"] = [
                    {
                        "name": sheet["name"],
                        "role": sheet["role"],
                        "headers": sheet["headers"],
                        "row_count": sheet["row_count"],
                        "column_count": sheet["column_count"],
                        "data_range": sheet["data_range"],
                        "detected_concentrations": sheet["detected_concentrations"],
                    }
                    for sheet in structure["sheets"]
                ]
                metadata["concentrations"] = list(structure["concentrations"])
            else:
                if structure["error"]:
                    metadata["error"] = structure["error"]
                metadata["notes"] = structure["error"] or "workbook structure unavailable"

            self.logger.info(f"Extracted Excel metadata from {file_path}")

        except Exception as e:
            self.logger.error(f"Error extracting Excel metadata from {file_path}: {e}")
            metadata["error"] = str(e)

        return metadata

    def extract_flowjo_metadata(self, file_path: str) -> Dict[str, Any]:
        """
        Extract metadata from FlowJo workspace files.

        Args:
            file_path: Path to the WSP file

        Returns:
            Dictionary of extracted metadata
        """
        metadata = {
            "file_type": "flowjo_wsp",
            "extraction_method": "basic_file_info",
            "notes": "Full FlowJo metadata extraction requires specialized library",
        }

        try:
            # Basic file info
            stat_info = os.stat(file_path)
            metadata["file_size"] = str(stat_info.st_size)
            metadata["created_date"] = str(stat_info.st_ctime)
            metadata["modified_date"] = str(stat_info.st_mtime)

            # Read file as text (WSP files are XML-based)
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()

                # Look for common patterns
                if "<Workspace" in content:
                    metadata["format"] = "FlowJo Workspace"
                if "gating" in content.lower():
                    metadata["has_gating"] = "true"
                if "compensation" in content.lower():
                    metadata["has_compensation"] = "true"

            self.logger.info(f"Extracted FlowJo metadata from {file_path}")

        except Exception as e:
            self.logger.error(f"Error extracting FlowJo metadata from {file_path}: {e}")
            metadata["error"] = str(e)

        return metadata

    def extract_all_metadata(self, folder_path: str) -> ExperimentMetadata:
        """
        Extract all metadata from an experiment folder.

        Args:
            folder_path: Path to the experiment folder

        Returns:
            ExperimentMetadata object with all extracted metadata
        """
        folder = Path(folder_path)

        # Extract experiment ID and name from folder name
        experiment_id = folder.name.split("_")[0] if "_" in folder.name else folder.name
        experiment_name = folder.name

        metadata = ExperimentMetadata(
            experiment_id=experiment_id, experiment_name=experiment_name, folder_path=str(folder)
        )

        # Process all files
        for item in folder.rglob("*"):
            if item.is_file():
                try:
                    file_metadata = self._extract_file_metadata(item)
                    metadata.files.append(file_metadata)
                except Exception as e:
                    self.logger.error(f"Error processing {item}: {e}")

        self.logger.info(f"Extracted metadata for {len(metadata.files)} files from {experiment_id}")
        return metadata

    def _extract_file_metadata(self, file_path: Path) -> FileMetadata:
        """
        Extract metadata from a single file.

        Args:
            file_path: Path to the file

        Returns:
            FileMetadata object
        """
        suffix = file_path.suffix.lower()
        checksum = self.calculate_checksum(str(file_path))

        # Get basic file info
        stat_info = file_path.stat()

        file_metadata = FileMetadata(
            file_path=str(file_path),
            file_name=file_path.name,
            file_type=suffix,
            file_size=int(stat_info.st_size),
            created_date=str(stat_info.st_ctime),
            modified_date=str(stat_info.st_mtime),
            checksum=checksum,
        )

        # Extract type-specific metadata
        if suffix == ".czi":
            file_metadata.raw_metadata = self.extract_czi_metadata(str(file_path))
        elif suffix in [".tiff", ".tif"]:
            file_metadata.raw_metadata = self.extract_tiff_metadata(str(file_path))
        elif suffix == ".ndpi":
            file_metadata.raw_metadata = self._extract_ndpi_metadata(str(file_path))
        elif suffix == ".lif":
            file_metadata.raw_metadata = self._extract_lif_metadata(str(file_path))
        elif suffix == ".xml":
            if "tiff_metadata" in file_path.name:
                file_metadata.raw_metadata = self.extract_tiff_metadata_xml(str(file_path))
            else:
                file_metadata.raw_metadata = self.extract_xml_metadata(str(file_path))
        elif suffix in [".xlif", ".lof", ".xlef", ".lifext", ".xsl"]:
            # Leica LAS X metadata files are XML-based
            file_metadata.raw_metadata = self._extract_leica_metadata(str(file_path))
        elif suffix == ".pzfx":
            # GraphPad Prism files are XML-based
            file_metadata.raw_metadata = self._extract_prism_metadata(str(file_path))
        elif suffix == ".fcs":
            file_metadata.raw_metadata = self.extract_fcs_metadata(str(file_path))
        elif suffix == ".wsp":
            file_metadata.raw_metadata = self.extract_flowjo_metadata(str(file_path))
        elif suffix in [".xlsx", ".xlsm", ".xls"]:
            file_metadata.raw_metadata = self.extract_excel_metadata(str(file_path))
        else:
            file_metadata.raw_metadata = {
                "file_type": suffix,
                "notes": "No specialized metadata extraction for this file type",
            }

        return file_metadata

    # ------------------------------------------------------------------
    # New format-specific extractors for partner data
    # ------------------------------------------------------------------

    def _extract_ndpi_metadata(self, file_path: str) -> Dict[str, Any]:
        """Extract metadata from Hamamatsu NanoZoomer (.ndpi) slide scanner files.

        NDPI files are big-TIFF variants.  This extractor reads the basic
        TIFF header and, when ``tifffile`` is available, the NDPI-specific
        tags (objective magnification, physical pixel size, etc.).
        """
        metadata: Dict[str, Any] = {
            "file_type": "ndpi",
            "extraction_method": "basic_plus_tifffile",
        }
        try:
            stat_info = os.stat(file_path)
            metadata["file_size"] = str(stat_info.st_size)
            metadata["modified_date"] = str(stat_info.st_mtime)

            # Try tifffile for rich NDPI tag extraction
            try:
                import tifffile as _tifffile  # type: ignore

                with _tifffile.TiffFile(file_path) as tif:
                    page = tif.pages[0]
                    metadata["image_width"] = page.shape[1] if len(page.shape) > 1 else ""
                    metadata["image_height"] = page.shape[0] if len(page.shape) > 0 else ""
                    metadata["compression"] = str(page.compression)
                    # NDPI-specific tags
                    for tag_id, tag in getattr(page, "tags", {}).items():
                        try:
                            metadata[f"tag_{tag_id}"] = str(tag.value)[:200]
                        except Exception:
                            pass
            except ImportError:
                metadata["note"] = "tifffile not installed – only basic file info extracted"
            except Exception as e:
                metadata["note"] = f"tifffile read warning: {e}"

            self.logger.info(f"Extracted NDPI metadata from {file_path}")
        except Exception as e:
            self.logger.error(f"Error extracting NDPI metadata: {e}")
            metadata["error"] = str(e)
        return metadata

    def _extract_lif_metadata(self, file_path: str) -> Dict[str, Any]:
        """Extract metadata from Leica LAS X (.lif) confocal image files.

        Uses ``readlif`` when available to enumerate image series, dimensions,
        and channel names.  Falls back to basic file info.
        """
        metadata: Dict[str, Any] = {
            "file_type": "lif",
            "extraction_method": "basic_plus_readlif",
        }
        try:
            stat_info = os.stat(file_path)
            metadata["file_size"] = str(stat_info.st_size)
            metadata["modified_date"] = str(stat_info.st_mtime)

            try:
                from readlif.reader import LifFile  # type: ignore

                lif = LifFile(file_path)
                series_info: List[Dict[str, Any]] = []
                for img in lif.get_iter_image():
                    info: Dict[str, Any] = {"name": img.name}
                    try:
                        info["dims"] = {
                            "x": img.dims.x,
                            "y": img.dims.y,
                            "z": img.dims.z,
                            "t": img.dims.t,
                            "c": img.dims.c,
                        }
                    except Exception:
                        pass
                    series_info.append(info)
                metadata["series_count"] = len(series_info)
                metadata["series"] = series_info
            except ImportError:
                metadata["note"] = "readlif not installed – only basic file info extracted"
            except Exception as e:
                metadata["note"] = f"readlif read warning: {e}"

            self.logger.info(f"Extracted LIF metadata from {file_path}")
        except Exception as e:
            self.logger.error(f"Error extracting LIF metadata: {e}")
            metadata["error"] = str(e)
        return metadata

    def _extract_leica_metadata(self, file_path: str) -> Dict[str, Any]:
        """Extract metadata from Leica LAS X XML-based files (.xlif, .lof, .xlef, .lifext, .xsl).

        These are standard XML files and are parsed with the generic XML extractor.
        """
        return self.extract_xml_metadata(file_path)

    def _extract_prism_metadata(self, file_path: str) -> Dict[str, Any]:
        """Extract metadata from GraphPad Prism (.pzfx) project files.

        Prism files are XML-based and parsed with the generic XML extractor.
        """
        return self.extract_xml_metadata(file_path)

    def extract_xml_metadata(self, file_path: str) -> Dict[str, Any]:
        """
        Extract metadata from generic XML files.

        Args:
            file_path: Path to the XML file

        Returns:
            Dictionary of extracted metadata
        """
        metadata = {"file_type": "xml", "extraction_method": "xml_parsing"}

        try:
            tree = ET.parse(file_path)
            root = tree.getroot()

            # Extract all elements
            def extract_elements(element, prefix=""):
                result = {}
                for child in element:
                    key = f"{prefix}/{child.tag}" if prefix else child.tag
                    if len(child) > 0:
                        result[key] = extract_elements(child, key)
                    else:
                        result[key] = child.text
                return result

            metadata["xml_data"] = json.dumps(extract_elements(root), default=str)

            self.logger.info(f"Extracted XML metadata from {file_path}")

        except Exception as e:
            self.logger.error(f"Error extracting XML metadata from {file_path}: {e}")
            metadata["error"] = str(e)

        return metadata

    def preserve_original_metadata(self, source_path: str, output_path: str):
        """
        Preserve original metadata by creating a sidecar JSON file.

        Args:
            source_path: Path to the original file
            output_path: Path where metadata should be saved
        """
        file_path = Path(source_path)
        metadata_path = Path(output_path)

        try:
            # Extract metadata
            if file_path.suffix == ".czi":
                metadata = self.extract_czi_metadata(str(file_path))
            elif file_path.suffix == ".tiff" or file_path.suffix == ".tif":
                metadata = self.extract_tiff_metadata(str(file_path))
            elif file_path.suffix == ".xml":
                if "tiff_metadata" in file_path.name:
                    metadata = self.extract_tiff_metadata_xml(str(file_path))
                else:
                    metadata = self.extract_xml_metadata(str(file_path))
            elif file_path.suffix == ".fcs":
                metadata = self.extract_fcs_metadata(str(file_path))
            elif file_path.suffix == ".wsp":
                metadata = self.extract_flowjo_metadata(str(file_path))
            elif file_path.suffix in [".xlsx", ".xlsm", ".xls"]:
                metadata = self.extract_excel_metadata(str(file_path))
            else:
                metadata = {
                    "file_type": file_path.suffix,
                    "notes": "No specialized metadata extraction for this file type",
                }

            # Add file info
            metadata["original_file"] = str(file_path)
            metadata["original_name"] = file_path.name
            metadata["original_size"] = str(file_path.stat().st_size)
            metadata["original_modified"] = str(file_path.stat().st_mtime)

            # Save as JSON
            metadata_path.parent.mkdir(parents=True, exist_ok=True)
            with open(metadata_path, "w", encoding="utf-8") as f:
                json.dump(metadata, f, indent=2, default=str, ensure_ascii=False)

            self.logger.info(f"Preserved metadata for {file_path.name} to {metadata_path}")

        except Exception as e:
            self.logger.error(f"Error preserving metadata for {source_path}: {e}")


def main():
    """Main function for testing the metadata extractor."""
    import sys

    # Test with representative data folder
    test_folder = "partner representative data/E1_Müller_Calceinassay und FACS Test"

    if len(sys.argv) > 1:
        test_folder = sys.argv[1]

    extractor = MetadataExtractor()
    metadata = extractor.extract_all_metadata(test_folder)

    print(f"\nMetadata extraction for: {metadata.experiment_name}")
    print(f"Files processed: {len(metadata.files)}")
    print("\nFile types:")
    for file_meta in metadata.files:
        print(f"  {file_meta.file_name} ({file_meta.file_type}): {file_meta.file_size} bytes")


if __name__ == "__main__":
    main()
