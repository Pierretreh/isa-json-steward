"""
Tests for the microscopy .docx parsing script (scripts/parse_microscopy_docx.py).

These tests use synthetic .docx files and study JSON structures in temporary
directories to avoid requiring the actual microscopy data.
"""

import json
import os
import sys
from pathlib import Path

import pytest

# Discover profile scripts directory
_profile_dir = None
_env = os.environ.get("ISA_STEWARD_PROFILE")
if _env and Path(_env).is_dir():
    _profile_dir = Path(_env)
else:
    _candidates = sorted(Path(__file__).parent.parent.glob("*-profile"))
    if _candidates:
        _profile_dir = _candidates[0]

if _profile_dir and (_profile_dir / "scripts").is_dir():
    sys.path.insert(0, str(_profile_dir / "scripts"))
try:
    from parse_microscopy_docx import (
        _make_observation_comment,
        add_observations_to_study,
        assess_aggregation,
        assess_llps,
        build_observations,
        extract_concentrations,
        extract_conditions,
        extract_docx_text,
        extract_protein_variants,
        process_docx_for_mapping,
    )
except ImportError:
    pytest.skip(
        "parse_microscopy_docx not available (requires profile with scripts/)",
        allow_module_level=True,
    )

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
def docx_with_llps(tmp_path: Path) -> Path:
    """Create a synthetic .docx file with LLPS-positive content."""
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx not installed")

    doc = Document()
    doc.add_heading("Microscopy Observations", level=1)
    doc.add_paragraph("2.12.25 (vitreous humor buffer without CaCl2)")
    doc.add_paragraph("C4(1) 2.25 uM pVV018 + 1.8 uM Eylea + 1 uM PPEP1")
    doc.add_paragraph(
        "After 1 hour at 37°C, clear droplets visible indicating "
        "liquid-liquid phase separation (LLPS). Spherical condensates "
        "observed. No aggregation detected."
    )
    doc.add_paragraph("C5(2) 2.25 uM pVV019 + 1.8 uM Eylea + 1 uM PPEP1")
    doc.add_paragraph("Droplets visible. Phase separation confirmed.")

    # Add a table.
    table = doc.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Well"
    table.rows[0].cells[1].text = "Sample"
    table.rows[0].cells[2].text = "Observation"
    table.rows[1].cells[0].text = "C4"
    table.rows[1].cells[1].text = "pVV018"
    table.rows[1].cells[2].text = "LLPS visible, round droplets"

    path = tmp_path / "observations_llps.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def docx_with_aggregation(tmp_path: Path) -> Path:
    """Create a synthetic .docx file with aggregation content."""
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx not installed")

    doc = Document()
    doc.add_paragraph("Microscopy 27.3.25")
    doc.add_paragraph(
        "25 uM pVV014 + 50 uM PPEP1. After cleavage at 37°C, "
        "heavy aggregation observed. Turbid solution with visible "
        "precipitate and fibril formation. Sediment at bottom."
    )
    doc.add_paragraph("20 uM pVV014 + 40 uM PPEP1. Similar aggregation pattern.")

    path = tmp_path / "observations_agg.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def docx_empty(tmp_path: Path) -> Path:
    """Create an empty .docx file."""
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx not installed")

    doc = Document()
    path = tmp_path / "empty.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def docx_no_observations(tmp_path: Path) -> Path:
    """Create a .docx file with only well plate layout (no observations)."""
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx not installed")

    doc = Document()
    doc.add_paragraph("384 well plate layout")
    doc.add_paragraph("D4 (1) MBP pVV002 42 uM as a control")
    doc.add_paragraph("D6 (2) 42 uM pVV014 + 7 uM PPEP1")
    doc.add_paragraph("F4 (3) MBP pVV002 40 uM as a control")

    path = tmp_path / "layout_only.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def study_dir(tmp_path: Path) -> Path:
    """Create a synthetic study directory with a study.json."""
    studies_root = tmp_path / "investigations" / "inv_inm" / "studies"
    study = studies_root / "study_SD24105_Analysis_of_018-024"
    study.mkdir(parents=True)
    (study / "files" / "original").mkdir(parents=True)

    study_json = {
        "studies": [
            {
                "title": "Analysis of 018-024 (SD24105)",
                "comments": [
                    {"name": "SD Number", "value": "SD24105"},
                ],
                "assays": [
                    {
                        "@id": "#assay_Depot_microscopy_assay",
                        "name": "Depot microscopy assay",
                        "dataFiles": [],
                        "materials": {"samples": [], "otherMaterials": []},
                        "processSequence": [],
                    },
                ],
            }
        ]
    }
    (study / "study.json").write_text(json.dumps(study_json, indent=2), encoding="utf-8")
    return study


@pytest.fixture
def microscopy_root(tmp_path: Path) -> Path:
    """Create a synthetic microscopy data directory."""
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx not installed")

    root = tmp_path / "Microscopy of LLPS"
    root.mkdir()

    # Create a mapped folder with a .docx file.
    folder = (
        root / "[ordered] Microscopy of 2.25 uM depot + Eylea" / "021225 018-024 eye buffer + eylea"
    )
    folder.mkdir(parents=True)

    doc = Document()
    doc.add_paragraph(
        "2.12.25 vitreous humor buffer. LLPS droplets visible "
        "after 1 hour at 37°C with PPEP1 cleavage. "
        "2.25 uM pVV018 + 1.8 uM Eylea + 1 uM PPEP1."
    )
    doc.save(str(folder / "18-24 2.12.25.docx"))

    return root


@pytest.fixture
def dm(tmp_path: Path):
    """Create a DirectoryManager rooted at tmp_path."""
    from utils.directory_manager import DirectoryManager

    return DirectoryManager(base_path=str(tmp_path))


# ---------------------------------------------------------------------------
# Tests: extract_docx_text
# ---------------------------------------------------------------------------


class TestExtractDocxText:
    """Tests for extract_docx_text()."""

    def test_extracts_paragraphs(self, docx_with_llps: Path):
        text = extract_docx_text(docx_with_llps)
        assert "Microscopy Observations" in text
        assert "pVV018" in text
        assert "liquid-liquid phase separation" in text

    def test_extracts_tables(self, docx_with_llps: Path):
        text = extract_docx_text(docx_with_llps)
        assert "Well" in text
        assert "pVV018" in text

    def test_empty_document(self, docx_empty: Path):
        text = extract_docx_text(docx_empty)
        assert text.strip() == ""

    def test_nonexistent_file(self, tmp_path: Path):
        text = extract_docx_text(tmp_path / "nonexistent.docx")
        assert text == ""


# ---------------------------------------------------------------------------
# Tests: assess_llps
# ---------------------------------------------------------------------------


class TestAssessLLPS:
    """Tests for assess_llps()."""

    def test_positive_llps(self):
        text = "Droplets visible. Liquid-liquid phase separation observed."
        result, evidence = assess_llps(text)
        assert result is True
        assert len(evidence) > 0

    def test_negative_llps(self):
        text = "No droplets. Clear, homogeneous solution. No phase separation."
        result, evidence = assess_llps(text)
        assert result is False

    def test_ambiguous(self):
        text = "Sample was observed under the microscope."
        result, evidence = assess_llps(text)
        assert result is None

    def test_case_insensitive(self):
        text = "LLPS visible with DROPLETS forming."
        result, evidence = assess_llps(text)
        assert result is True

    def test_mixed_signals(self):
        # More positive than negative.
        text = "Droplets visible but also some clear areas. Phase separation confirmed."
        result, evidence = assess_llps(text)
        assert result is True


# ---------------------------------------------------------------------------
# Tests: assess_aggregation
# ---------------------------------------------------------------------------


class TestAssessAggregation:
    """Tests for assess_aggregation()."""

    def test_positive_aggregation(self):
        text = "Heavy aggregation observed. Turbid with precipitate."
        result, evidence = assess_aggregation(text)
        assert result is True

    def test_negative_aggregation(self):
        text = "No aggregation. Clear solution, soluble protein."
        result, evidence = assess_aggregation(text)
        assert result is False

    def test_ambiguous(self):
        text = "Sample observed at 37°C."
        result, evidence = assess_aggregation(text)
        assert result is None


# ---------------------------------------------------------------------------
# Tests: extract_conditions
# ---------------------------------------------------------------------------


class TestExtractConditions:
    """Tests for extract_conditions()."""

    def test_cleavage_detected(self):
        text = "After PPEP1 cleavage at 37°C"
        conds = extract_conditions(text)
        assert conds["cleavage"] == "yes"
        assert conds["ppep1_present"] == "yes"
        assert conds["temperature"] == "37°C"

    def test_no_cleavage(self):
        text = "Sample observed without treatment."
        conds = extract_conditions(text)
        assert conds["cleavage"] is None
        assert conds["ppep1_present"] is None

    def test_four_degrees(self):
        text = "O/N cleavage at 4°C"
        conds = extract_conditions(text)
        assert conds["temperature"] == "4°C"


# ---------------------------------------------------------------------------
# Tests: extract_protein_variants
# ---------------------------------------------------------------------------


class TestExtractProteinVariants:
    """Tests for extract_protein_variants()."""

    def test_finds_variants(self):
        text = "2.25 uM pVV018 + 1.8 uM Eylea + 1 uM PPEP1"
        variants = extract_protein_variants(text)
        assert "PVV018" in variants

    def test_multiple_variants(self):
        text = "pVV018 and pVV019 and pVV021 tested"
        variants = extract_protein_variants(text)
        assert len(variants) == 3
        assert "PVV018" in variants
        assert "PVV019" in variants
        assert "PVV021" in variants

    def test_case_insensitive(self):
        text = "pVV014 and PVV014"
        variants = extract_protein_variants(text)
        assert len(variants) == 1  # deduplicated

    def test_no_variants(self):
        text = "Just some random text"
        variants = extract_protein_variants(text)
        assert variants == []


# ---------------------------------------------------------------------------
# Tests: extract_concentrations
# ---------------------------------------------------------------------------


class TestExtractConcentrations:
    """Tests for extract_concentrations()."""

    def test_finds_concentrations(self):
        text = "2.25 uM pVV018 + 1.8 uM Eylea + 1 uM PPEP1"
        concs = extract_concentrations(text)
        assert "2.25 uM" in concs
        assert "1.8 uM" in concs
        assert "1 uM" in concs

    def test_deduplication(self):
        text = "40 uM protein and 40 uM PPEP1"
        concs = extract_concentrations(text)
        assert concs.count("40 uM") == 1

    def test_no_concentrations(self):
        text = "No concentrations mentioned"
        concs = extract_concentrations(text)
        assert concs == []


# ---------------------------------------------------------------------------
# Tests: build_observations
# ---------------------------------------------------------------------------


class TestBuildObservations:
    """Tests for build_observations()."""

    def test_llps_observation(self, docx_with_llps: Path):
        text = extract_docx_text(docx_with_llps)
        obs = build_observations(docx_with_llps, text)
        assert obs["llps"]["present"] is True
        assert obs["source_file"] == docx_with_llps.name
        assert len(obs["protein_variants"]) > 0

    def test_aggregation_observation(self, docx_with_aggregation: Path):
        text = extract_docx_text(docx_with_aggregation)
        obs = build_observations(docx_with_aggregation, text)
        assert obs["aggregation"]["present"] is True

    def test_conditions_extracted(self, docx_with_llps: Path):
        text = extract_docx_text(docx_with_llps)
        obs = build_observations(docx_with_llps, text)
        assert obs["conditions"]["ppep1_present"] == "yes"


# ---------------------------------------------------------------------------
# Tests: _make_observation_comment
# ---------------------------------------------------------------------------


class TestMakeObservationComment:
    """Tests for _make_observation_comment()."""

    def test_comment_structure(self, docx_with_llps: Path):
        text = extract_docx_text(docx_with_llps)
        obs = build_observations(docx_with_llps, text)
        comment = _make_observation_comment(obs)
        assert "name" in comment
        assert "value" in comment
        assert "Microscopy Notes" in comment["name"]
        assert "LLPS: yes" in comment["value"]

    def test_aggregation_comment(self, docx_with_aggregation: Path):
        text = extract_docx_text(docx_with_aggregation)
        obs = build_observations(docx_with_aggregation, text)
        comment = _make_observation_comment(obs)
        assert "Aggregation: yes" in comment["value"]


# ---------------------------------------------------------------------------
# Tests: add_observations_to_study
# ---------------------------------------------------------------------------


class TestAddObservationsToStudy:
    """Tests for add_observations_to_study()."""

    def test_adds_comment(self, docx_with_llps: Path, study_dir: Path):
        json_path = study_dir / "study.json"
        text = extract_docx_text(docx_with_llps)
        obs = build_observations(docx_with_llps, text)
        added = add_observations_to_study(json_path, obs)
        assert added is True

        with open(json_path, encoding="utf-8") as f:
            data = json.load(f)
        comments = data["studies"][0]["comments"]
        note_comments = [c for c in comments if "Microscopy Notes" in c["name"]]
        assert len(note_comments) == 1
        assert "LLPS" in note_comments[0]["value"]

    def test_skips_duplicate(self, docx_with_llps: Path, study_dir: Path):
        json_path = study_dir / "study.json"
        text = extract_docx_text(docx_with_llps)
        obs = build_observations(docx_with_llps, text)
        add_observations_to_study(json_path, obs)
        added_again = add_observations_to_study(json_path, obs)
        assert added_again is False

    def test_dry_run_no_write(self, docx_with_llps: Path, study_dir: Path):
        json_path = study_dir / "study.json"
        text = extract_docx_text(docx_with_llps)
        obs = build_observations(docx_with_llps, text)
        added = add_observations_to_study(json_path, obs, dry_run=True)
        assert added is True

        with open(json_path, encoding="utf-8") as f:
            data = json.load(f)
        comments = data["studies"][0]["comments"]
        note_comments = [c for c in comments if "Microscopy Notes" in c["name"]]
        assert len(note_comments) == 0


# ---------------------------------------------------------------------------
# Tests: process_docx_for_mapping (integration)
# ---------------------------------------------------------------------------


class TestProcessDocxForMapping:
    """Integration tests for process_docx_for_mapping()."""

    def test_full_pipeline(
        self,
        microscopy_root: Path,
        study_dir: Path,
        dm,
        monkeypatch,
    ):
        monkeypatch.setattr("parse_microscopy_docx.MICROSCOPY_ROOT", microscopy_root)
        summary = process_docx_for_mapping(
            "[ordered] Microscopy of 2.25 uM depot + Eylea/021225 018-024 eye buffer + eylea",
            "SD24105",
            dm,
        )
        assert summary["sd_number"] == "SD24105"
        assert summary["docx_found"] >= 1
        assert summary["comments_added"] >= 1
        assert summary["errors"] == []

    def test_dry_run(
        self,
        microscopy_root: Path,
        study_dir: Path,
        dm,
        monkeypatch,
    ):
        monkeypatch.setattr("parse_microscopy_docx.MICROSCOPY_ROOT", microscopy_root)
        summary = process_docx_for_mapping(
            "[ordered] Microscopy of 2.25 uM depot + Eylea/021225 018-024 eye buffer + eylea",
            "SD24105",
            dm,
            dry_run=True,
        )
        assert summary["docx_found"] >= 1
        assert summary["comments_added"] >= 1

        # Verify nothing was written.
        json_path = study_dir / "study.json"
        with open(json_path, encoding="utf-8") as f:
            data = json.load(f)
        comments = data["studies"][0]["comments"]
        note_comments = [c for c in comments if "Microscopy Notes" in c["name"]]
        assert len(note_comments) == 0

    def test_missing_source(
        self,
        study_dir: Path,
        dm,
        tmp_path: Path,
        monkeypatch,
    ):
        empty_root = tmp_path / "empty_root"
        empty_root.mkdir()
        monkeypatch.setattr("parse_microscopy_docx.MICROSCOPY_ROOT", empty_root)
        summary = process_docx_for_mapping(
            "nonexistent/path",
            "SD24105",
            dm,
        )
        assert summary["errors"]
        assert "not found" in summary["errors"][0].lower()

    def test_missing_study(
        self,
        microscopy_root: Path,
        dm,
        monkeypatch,
    ):
        monkeypatch.setattr("parse_microscopy_docx.MICROSCOPY_ROOT", microscopy_root)
        summary = process_docx_for_mapping(
            "[ordered] Microscopy of 2.25 uM depot + Eylea/021225 018-024 eye buffer + eylea",
            "SD99999",
            dm,
        )
        assert summary["errors"]
        assert "study" in summary["errors"][0].lower()


# ---------------------------------------------------------------------------
# Tests: constants
# ---------------------------------------------------------------------------


class TestConstants:
    """Sanity checks on module-level constants."""

    def test_llps_keywords_not_empty(self):
        from parse_microscopy_docx import LLPS_POSITIVE_KEYWORDS

        assert len(LLPS_POSITIVE_KEYWORDS) > 0

    def test_aggregation_keywords_not_empty(self):
        from parse_microscopy_docx import AGGREGATION_POSITIVE_KEYWORDS

        assert len(AGGREGATION_POSITIVE_KEYWORDS) > 0

    def test_cleavage_keywords_not_empty(self):
        from parse_microscopy_docx import CLEAVAGE_KEYWORDS

        assert len(CLEAVAGE_KEYWORDS) > 0
