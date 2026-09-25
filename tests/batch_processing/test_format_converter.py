"""Tests for FormatConverter in utils/batch/format_converter.py.

Covers:
  - ConversionResult dataclass
  - FormatConverter.__init__
  - _parse_fcs_header (static)
  - _parse_fcs_keywords (static)
  - _fcs_byte_order (static)
  - convert_fcs_to_csv
  - convert_excel_to_csv
  - convert_xml_to_json
  - get_conversion_manifest
  - batch_convert_folder
"""

import json
import struct
from pathlib import Path

import pytest

from utils.batch.format_converter import ConversionResult, FormatConverter

# ---------------------------------------------------------------------------
# Helpers to craft minimal FCS binary data
# ---------------------------------------------------------------------------


def _build_fcs_header(
    text_start: int,
    text_end: int,
    data_start: int = 0,
    data_end: int = 0,
    analysis_start: int = 0,
    analysis_end: int = 0,
    version: str = "FCS3.0",
) -> bytes:
    """Build a 58-byte FCS header with the given segment offsets."""
    hdr = bytearray(58)
    hdr[0:6] = version.encode("ascii").ljust(6)
    hdr[6:10] = b"    "
    hdr[10:18] = str(text_start).rjust(8).encode("ascii")
    hdr[18:26] = str(text_end).rjust(8).encode("ascii")
    hdr[26:34] = str(data_start).rjust(8).encode("ascii")
    hdr[34:42] = str(data_end).rjust(8).encode("ascii")
    hdr[42:50] = str(analysis_start).rjust(8).encode("ascii")
    hdr[50:58] = str(analysis_end).rjust(8).encode("ascii")
    return bytes(hdr)


def _build_fcs_text_segment(keywords: dict, delimiter: str = "/") -> bytes:
    """Build an FCS TEXT segment from a keyword dict.

    Format: /key1/value1/key2/value2/  (delimiter-separated key/value pairs)
    """
    parts = [""]
    for k, v in keywords.items():
        parts.extend([k, v])
    parts.append("")
    return delimiter.join(parts).encode("latin1")


def _build_minimal_fcs(keywords: dict = None, data_bytes: bytes = b"") -> bytes:
    """Build a minimal but valid FCS file as bytes."""
    if keywords is None:
        keywords = {
            "$PAR": "2",
            "$TOT": "2",
            "$P1N": "FSC-A",
            "$P2N": "SSC-A",
            "$BYTEORD": "1,2,3,4",
            "$DATATYPE": "I",
            "$MODE": "L",
            "$P1B": "16",
            "$P2B": "16",
        }

    text_segment = _build_fcs_text_segment(keywords)
    text_start = 58
    text_end = text_start + len(text_segment)

    # Data segment follows text segment
    if data_bytes:
        data_start = text_end
        data_end = data_start + len(data_bytes)
    else:
        data_start = 0
        data_end = 0

    header = _build_fcs_header(text_start, text_end, data_start, data_end)
    return header + text_segment + data_bytes


# ---------------------------------------------------------------------------
# ConversionResult dataclass tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestConversionResult:
    """Tests for the ConversionResult dataclass."""

    def test_creation_minimal(self):
        r = ConversionResult(
            source_file="a.fcs", target_file="a.csv", success=True, conversion_type="fcs_to_csv"
        )
        assert r.source_file == "a.fcs"
        assert r.target_file == "a.csv"
        assert r.success is True
        assert r.conversion_type == "fcs_to_csv"
        assert r.metadata_file is None
        assert r.error_message is None
        assert r.warnings == []

    def test_creation_with_all_fields(self):
        r = ConversionResult(
            source_file="a.fcs",
            target_file="a.csv",
            success=False,
            conversion_type="fcs_to_csv",
            metadata_file="a_meta.json",
            error_message="bad file",
            warnings=["warn1", "warn2"],
        )
        assert r.metadata_file == "a_meta.json"
        assert r.error_message == "bad file"
        assert r.warnings == ["warn1", "warn2"]


# ---------------------------------------------------------------------------
# FormatConverter.__init__ tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestFormatConverterInit:
    """Tests for FormatConverter initialization."""

    def test_default_init(self):
        fc = FormatConverter()
        assert fc.preserve_originals is True
        assert fc.metadata_extractor is not None

    def test_init_no_preserve(self):
        fc = FormatConverter(preserve_originals=False)
        assert fc.preserve_originals is False


# ---------------------------------------------------------------------------
# _parse_fcs_header tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestParseFcsHeader:
    """Tests for FormatConverter._parse_fcs_header static method."""

    def test_valid_fcs30_header(self):
        raw = _build_fcs_header(58, 200, 200, 400, 0, 0)
        hdr = FormatConverter._parse_fcs_header(raw)
        assert hdr["version"] == "FCS3.0"
        assert hdr["text_start"] == 58
        assert hdr["text_end"] == 200
        assert hdr["data_start"] == 200
        assert hdr["data_end"] == 400
        assert hdr["analysis_start"] == 0
        assert hdr["analysis_end"] == 0

    def test_valid_fcs31_header(self):
        raw = _build_fcs_header(100, 300, version="FCS3.1")
        hdr = FormatConverter._parse_fcs_header(raw)
        assert hdr["version"] == "FCS3.1"
        assert hdr["text_start"] == 100
        assert hdr["text_end"] == 300

    def test_too_small_raises(self):
        with pytest.raises(ValueError, match="too small"):
            FormatConverter._parse_fcs_header(b"tiny")

    def test_zero_offsets(self):
        raw = _build_fcs_header(0, 0, 0, 0, 0, 0)
        hdr = FormatConverter._parse_fcs_header(raw)
        assert hdr["text_start"] == 0
        assert hdr["data_start"] == 0


# ---------------------------------------------------------------------------
# _parse_fcs_keywords tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestParseFcsKeywords:
    """Tests for FormatConverter._parse_fcs_keywords static method."""

    def test_valid_keywords(self):
        raw = _build_minimal_fcs()
        # Text segment starts at byte 58
        text_start = 58
        text_end = len(raw)
        kw = FormatConverter._parse_fcs_keywords(raw, text_start, text_end)
        assert kw["$PAR"] == "2"
        assert kw["$TOT"] == "2"
        assert kw["$P1N"] == "FSC-A"
        assert kw["$P2N"] == "SSC-A"

    def test_empty_text_segment(self):
        raw = _build_fcs_header(0, 0)
        kw = FormatConverter._parse_fcs_keywords(raw, 0, 0)
        assert kw == {}

    def test_invalid_offsets(self):
        raw = _build_minimal_fcs()
        # text_start > text_end
        kw = FormatConverter._parse_fcs_keywords(raw, 200, 58)
        assert kw == {}

    def test_text_end_exceeds_raw(self):
        raw = _build_minimal_fcs()
        kw = FormatConverter._parse_fcs_keywords(raw, 58, 99999)
        assert kw == {}

    def test_custom_delimiter(self):
        # Build text segment with '|' delimiter
        text = "|$PAR|3|$TOT|50|".encode("latin1")
        header = _build_fcs_header(58, 58 + len(text))
        raw = header + text
        kw = FormatConverter._parse_fcs_keywords(raw, 58, 58 + len(text))
        assert kw["$PAR"] == "3"
        assert kw["$TOT"] == "50"


# ---------------------------------------------------------------------------
# _fcs_byte_order tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestFcsByteOrder:
    """Tests for FormatConverter._fcs_byte_order static method."""

    def test_little_endian(self):
        assert FormatConverter._fcs_byte_order({"$BYTEORD": "1,2,3,4"}) == "little"

    def test_big_endian(self):
        assert FormatConverter._fcs_byte_order({"$BYTEORD": "4,3,2,1"}) == "big"

    def test_default_little(self):
        assert FormatConverter._fcs_byte_order({}) == "little"


# ---------------------------------------------------------------------------
# convert_fcs_to_csv tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestConvertFcsToCsv:
    """Tests for FormatConverter.convert_fcs_to_csv."""

    def test_convert_fcs_with_integer_data(self, tmp_path):
        """Test FCS→CSV conversion with integer data segment."""
        # Build a minimal FCS file with 2 events, 2 parameters, 16-bit integers
        keywords = {
            "$PAR": "2",
            "$TOT": "2",
            "$P1N": "FSC-A",
            "$P2N": "SSC-A",
            "$BYTEORD": "1,2,3,4",
            "$DATATYPE": "I",
            "$MODE": "L",
            "$P1B": "16",
            "$P2B": "16",
            "$DATE": "01-JAN-2024",
            "$CYT": "TestCyt",
        }
        # 2 events × 2 params × 2 bytes = 8 bytes of data
        data = struct.pack("<HHHH", 100, 200, 300, 400)
        fcs_bytes = _build_minimal_fcs(keywords, data)

        fcs_path = tmp_path / "test.fcs"
        fcs_path.write_bytes(fcs_bytes)
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        csv_path, meta_path = fc.convert_fcs_to_csv(str(fcs_path), str(output_dir))

        assert csv_path != ""
        assert meta_path != ""
        assert Path(csv_path).exists()
        assert Path(meta_path).exists()

        # Verify CSV content
        content = Path(csv_path).read_text()
        assert "FSC-A" in content
        assert "SSC-A" in content

        # Verify metadata JSON
        with open(meta_path) as f:
            meta = json.load(f)
        assert meta["conversion"]["events_written"] == 2
        assert "FSC-A" in meta["conversion"]["parameters"]

    def test_convert_fcs_with_float_data(self, tmp_path):
        """Test FCS→CSV conversion with float data segment."""
        keywords = {
            "$PAR": "2",
            "$TOT": "1",
            "$P1N": "FITC-A",
            "$P2N": "PE-A",
            "$BYTEORD": "1,2,3,4",
            "$DATATYPE": "F",
            "$MODE": "L",
        }
        data = struct.pack("<ff", 1.5, 2.5)
        fcs_bytes = _build_minimal_fcs(keywords, data)

        fcs_path = tmp_path / "float.fcs"
        fcs_path.write_bytes(fcs_bytes)
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        csv_path, meta_path = fc.convert_fcs_to_csv(str(fcs_path), str(output_dir))
        assert csv_path != ""
        content = Path(csv_path).read_text()
        assert "FITC-A" in content

    def test_convert_fcs_no_data_segment(self, tmp_path):
        """Test FCS→CSV when DATA segment is empty (metadata-only CSV)."""
        keywords = {
            "$PAR": "2",
            "$TOT": "0",
            "$P1N": "FSC-A",
            "$P2N": "SSC-A",
            "$BYTEORD": "1,2,3,4",
            "$DATATYPE": "I",
            "$MODE": "L",
        }
        fcs_bytes = _build_minimal_fcs(keywords)

        fcs_path = tmp_path / "nodata.fcs"
        fcs_path.write_bytes(fcs_bytes)
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        csv_path, meta_path = fc.convert_fcs_to_csv(str(fcs_path), str(output_dir))
        assert csv_path != ""
        content = Path(csv_path).read_text()
        assert "Note" in content  # fallback row

    def test_convert_fcs_nonexistent_file(self, tmp_path):
        """Test convert_fcs_to_csv with nonexistent file returns empty strings."""
        fc = FormatConverter()
        csv_path, meta_path = fc.convert_fcs_to_csv(
            str(tmp_path / "nonexistent.fcs"), str(tmp_path / "out")
        )
        assert csv_path == ""
        assert meta_path == ""


# ---------------------------------------------------------------------------
# convert_excel_to_csv tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestConvertExcelToCsv:
    """Tests for FormatConverter.convert_excel_to_csv (real XLSX→CSV)."""

    # -- workbook fixtures (built with openpyxl, like test_parse_excel_metadata) --

    @staticmethod
    def _single_sheet_xlsx(tmp_path: Path) -> Path:
        import openpyxl

        path = tmp_path / "single.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["Sample", "Concentration (ng/ul)", "Readout"])
        ws.append(["S1", 0.5, 12.3])
        ws.append(["S2", 1.25, 45.6])
        ws.append(["S3", "5 ng/ml", 7])
        wb.save(str(path))
        wb.close()
        return path

    @staticmethod
    def _multi_sheet_xlsx(tmp_path: Path) -> Path:
        import openpyxl

        path = tmp_path / "multi.xlsx"
        wb = openpyxl.Workbook()
        notes = wb.active
        notes.title = "Notes"
        notes.append(["method", "value"])
        notes.append(["note one", "x"])
        data = wb.create_sheet("Data")
        data.append(["Sample", "Value"])
        data.append(["A", 1])
        data.append(["B", 2])
        data.append(["C", 3])
        wb.save(str(path))
        wb.close()
        return path

    @staticmethod
    def _read_sidecar(meta_path: str) -> dict:
        with open(meta_path, encoding="utf-8") as f:
            return json.load(f)

    def test_convert_excel_real_single_sheet(self, tmp_path):
        """Real conversion: CSV rows + sidecar schema (plan §2.1)."""
        xlsx_path = self._single_sheet_xlsx(tmp_path)
        fc = FormatConverter()
        csv_path, meta_path = fc.convert_excel_to_csv(str(xlsx_path), str(tmp_path / "out"))

        assert csv_path != ""
        assert Path(csv_path).exists()
        assert Path(meta_path).exists()

        rows = Path(csv_path).read_text(encoding="utf-8").splitlines()
        assert rows[0] == "Sample,Concentration (ng/ul),Readout"
        assert rows[1] == "S1,0.5,12.3"
        assert rows[2] == "S2,1.25,45.6"
        assert rows[3] == "S3,5 ng/ml,7"

        md = self._read_sidecar(meta_path)
        assert md["file_type"] == "excel"
        assert md["format"] == "XLSX (ZIP-based)"
        assert md["sheet_count"] == 1
        assert md["sheet_names"] == ["Sheet1"]
        assert md["conversion"]["real_conversion"] is True
        assert md["conversion"]["library"] == "openpyxl"
        assert md["conversion"]["primary_sheet"] == "Sheet1"
        assert md["conversion"]["rows_written"] == 3

        sheet = md["sheets"][0]
        assert sheet["is_primary"] is True
        assert sheet["role"] == "data"
        assert sheet["headers"] == ["Sample", "Concentration (ng/ul)", "Readout"]
        assert sheet["row_count"] == 3
        assert sheet["column_count"] == 3
        assert sheet["data_range"] == {
            "first_row": 1,
            "last_row": 4,
            "first_col": 1,
            "last_col": 3,
        }

    def test_convert_excel_multi_sheet_primary_detection(self, tmp_path):
        """Multi-sheet workbook: the 'Data' sheet is exported, not 'Notes'."""
        xlsx_path = self._multi_sheet_xlsx(tmp_path)
        fc = FormatConverter()
        csv_path, meta_path = fc.convert_excel_to_csv(str(xlsx_path), str(tmp_path / "out"))

        rows = Path(csv_path).read_text(encoding="utf-8").splitlines()
        assert rows[0] == "Sample,Value"
        assert len(rows) == 4  # header + 3 data rows

        md = self._read_sidecar(meta_path)
        assert md["sheet_names"] == ["Notes", "Data"]
        assert [s["is_primary"] for s in md["sheets"]] == [False, True]
        assert md["conversion"]["primary_sheet"] == "Data"
        assert md["conversion"]["real_conversion"] is True

    def test_convert_excel_formula_cell_value_not_formula(self, tmp_path):
        """Formula cells must never leak the formula string into the CSV."""
        import openpyxl

        path = tmp_path / "formula.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Calc"
        ws.append(["A", "B", "Sum"])
        ws.append([1, 2, "=A2+B2"])
        wb.save(str(path))
        wb.close()

        fc = FormatConverter()
        csv_path, meta_path = fc.convert_excel_to_csv(str(path), str(tmp_path / "out"))

        content = Path(csv_path).read_text(encoding="utf-8")
        assert "=A2+B2" not in content
        assert content.splitlines()[0] == "A,B,Sum"

        # openpyxl (written without Excel) has no cached value → empty cell +
        # a conversion note (plan §3.5)
        md = self._read_sidecar(meta_path)
        assert md["conversion"]["real_conversion"] is True
        assert "formula" in md["conversion"]["note"].lower()

    def test_convert_excel_merged_cells(self, tmp_path):
        """Merged cells: non-top-left cells come out empty."""
        import openpyxl

        path = tmp_path / "merged.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "M"
        ws.append(["Title", None])
        ws.append(["a", "b"])
        ws.merge_cells("A1:B1")
        wb.save(str(path))
        wb.close()

        fc = FormatConverter()
        csv_path, meta_path = fc.convert_excel_to_csv(str(path), str(tmp_path / "out"))
        rows = Path(csv_path).read_text(encoding="utf-8").splitlines()
        assert rows[0] == "Title,"
        assert rows[1] == "a,b"

    def test_convert_excel_concentration_detection_and_precision(self, tmp_path):
        """Concentration column detected; 0.5 kept exactly in the CSV."""
        xlsx_path = self._single_sheet_xlsx(tmp_path)
        fc = FormatConverter()
        csv_path, meta_path = fc.convert_excel_to_csv(str(xlsx_path), str(tmp_path / "out"))

        content = Path(csv_path).read_text(encoding="utf-8")
        assert "0.5" in content
        assert "0.50" not in content

        md = self._read_sidecar(meta_path)
        concs = md["sheets"][0]["detected_concentrations"]
        assert concs[0] == {
            "header": "Concentration (ng/ul)",
            "value": 0.5,
            "unit": "ng/ul",
            "row": 1,
        }
        assert concs[2]["value"] == 5.0
        assert concs[2]["unit"] == "ng/ml"
        assert len(md["concentrations"]) == 3

    def test_convert_excel_empty_sheet_excluded_from_primary(self, tmp_path):
        """An empty sheet first must not win primary-sheet selection."""
        import openpyxl

        path = tmp_path / "empty_first.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Empty"
        results = wb.create_sheet("Results")
        results.append(["x", "y"])
        results.append([1, 2])
        results.append([3, 4])
        wb.save(str(path))
        wb.close()

        fc = FormatConverter()
        csv_path, meta_path = fc.convert_excel_to_csv(str(path), str(tmp_path / "out"))

        rows = Path(csv_path).read_text(encoding="utf-8").splitlines()
        assert rows[0] == "x,y"
        assert len(rows) == 3  # header + 2 data rows

        md = self._read_sidecar(meta_path)
        assert md["conversion"]["primary_sheet"] == "Results"
        assert [s["is_primary"] for s in md["sheets"]] == [False, True]

    def test_convert_excel_non_ascii_headers(self, tmp_path):
        """Non-ASCII headers round-trip as UTF-8 in CSV and sidecar."""
        import openpyxl

        path = tmp_path / "nonascii.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Übersicht"
        ws.append(["Konzentration", "Messung"])
        ws.append(["5 ng/ul", 3.14])
        wb.save(str(path))
        wb.close()

        fc = FormatConverter()
        csv_path, meta_path = fc.convert_excel_to_csv(str(path), str(tmp_path / "out"))

        content = Path(csv_path).read_text(encoding="utf-8")
        assert "Konzentration" in content
        assert "5 ng/ul" in content

        md = self._read_sidecar(meta_path)
        assert md["sheets"][0]["headers"][0] == "Konzentration"
        assert md["concentrations"][0]["unit"] == "ng/ul"

    def test_convert_excel_degrades_on_invalid(self, tmp_path):
        """Corrupt file: CSV + sidecar still emitted, real_conversion=False."""
        xlsx_path = tmp_path / "data.xlsx"
        xlsx_path.write_bytes(b"\x50\x4b\x03\x04" + b"\x00" * 100)  # ZIP signature only
        fc = FormatConverter()
        csv_path, meta_path = fc.convert_excel_to_csv(str(xlsx_path), str(tmp_path / "out"))

        assert csv_path != ""
        assert Path(csv_path).exists()
        assert Path(meta_path).exists()

        rows = Path(csv_path).read_text(encoding="utf-8").splitlines()
        assert rows[0] == "Note"
        assert len(rows) == 2

        md = self._read_sidecar(meta_path)
        assert md["conversion"]["real_conversion"] is False
        assert md["conversion"]["note"]
        assert md["error"]
        assert md["sheet_count"] == 0

    def test_convert_excel_zero_byte_file(self, tmp_path):
        """Zero-byte file degrades with a 'zero-byte' error."""
        xlsx_path = tmp_path / "zero.xlsx"
        xlsx_path.write_bytes(b"")
        fc = FormatConverter()
        csv_path, meta_path = fc.convert_excel_to_csv(str(xlsx_path), str(tmp_path / "out"))

        assert csv_path != ""
        md = self._read_sidecar(meta_path)
        assert md["conversion"]["real_conversion"] is False
        assert "zero-byte" in md["conversion"]["note"]

    def test_convert_excel_legacy_xls_degrades(self, tmp_path):
        """Legacy .xls (OLE2) degrades with an 'export to .xlsx' hint."""
        xls_path = tmp_path / "legacy.xls"
        xls_path.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 100)  # OLE2 signature
        fc = FormatConverter()
        csv_path, meta_path = fc.convert_excel_to_csv(str(xls_path), str(tmp_path / "out"))

        assert csv_path != ""
        md = self._read_sidecar(meta_path)
        assert md["format"] == "XLS (OLE2)"
        assert md["conversion"]["real_conversion"] is False
        assert "export to .xlsx" in md["conversion"]["note"]

    def test_convert_excel_xlsm(self, tmp_path):
        """Macro-enabled .xlsm converts read-only, format 'XLSM (ZIP-based)'."""
        import openpyxl

        path = tmp_path / "macro.xlsm"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["a", "b"])
        ws.append([1, 2])
        wb.save(str(path))
        wb.close()

        fc = FormatConverter()
        csv_path, meta_path = fc.convert_excel_to_csv(str(path), str(tmp_path / "out"))

        rows = Path(csv_path).read_text(encoding="utf-8").splitlines()
        assert rows[0] == "a,b"
        assert rows[1] == "1,2"

        md = self._read_sidecar(meta_path)
        assert md["format"] == "XLSM (ZIP-based)"
        assert md["conversion"]["real_conversion"] is True
        assert md["conversion"]["primary_sheet"] == "Data"

    def test_convert_excel_nonexistent(self, tmp_path):
        """Missing input file → no output files at all ("", "")."""
        fc = FormatConverter()
        csv_path, meta_path = fc.convert_excel_to_csv(
            str(tmp_path / "missing.xlsx"), str(tmp_path / "out")
        )
        assert csv_path == ""
        assert meta_path == ""


# ---------------------------------------------------------------------------
# convert_xml_to_json tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestConvertXmlToJson:
    """Tests for FormatConverter.convert_xml_to_json."""

    def test_convert_xml_to_json(self, tmp_path):
        """Test XML→JSON conversion with a simple XML file."""
        xml_content = '<?xml version="1.0"?><root><item>value</item></root>'
        xml_path = tmp_path / "test.xml"
        xml_path.write_text(xml_content, encoding="utf-8")
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        json_path = fc.convert_xml_to_json(str(xml_path), str(output_dir))

        assert json_path != ""
        assert Path(json_path).exists()

        with open(json_path) as f:
            data = json.load(f)
        assert "item" in data

    def test_convert_xml_nonexistent(self, tmp_path):
        """Test convert_xml_to_json with nonexistent file."""
        fc = FormatConverter()
        json_path = fc.convert_xml_to_json(str(tmp_path / "missing.xml"), str(tmp_path / "out"))
        assert json_path == ""

    def test_convert_xml_nested(self, tmp_path):
        """Test XML→JSON with nested elements."""
        xml_content = '<?xml version="1.0"?><root><parent><child>val</child></parent></root>'
        xml_path = tmp_path / "nested.xml"
        xml_path.write_text(xml_content, encoding="utf-8")
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        json_path = fc.convert_xml_to_json(str(xml_path), str(output_dir))
        assert json_path != ""

        with open(json_path) as f:
            data = json.load(f)
        assert "parent" in data


# ---------------------------------------------------------------------------
# get_conversion_manifest tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestGetConversionManifest:
    """Tests for FormatConverter.get_conversion_manifest."""

    def test_manifest_exists(self, tmp_path):
        """Test reading an existing manifest."""
        manifest = {
            "conversion_date": "2024-01-01",
            "total_files": 5,
            "successful_conversions": 4,
        }
        manifest_path = tmp_path / "conversion_manifest.json"
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

        fc = FormatConverter()
        result = fc.get_conversion_manifest(str(tmp_path))
        assert result is not None
        assert result["total_files"] == 5

    def test_manifest_not_found(self, tmp_path):
        """Test when no manifest exists."""
        fc = FormatConverter()
        result = fc.get_conversion_manifest(str(tmp_path))
        assert result is None


# ---------------------------------------------------------------------------
# batch_convert_folder tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestBatchConvertFolder:
    """Tests for FormatConverter.batch_convert_folder."""

    def test_batch_convert_empty_folder(self, tmp_path):
        """Test batch conversion of an empty folder."""
        source = tmp_path / "source"
        source.mkdir()
        output = tmp_path / "output"

        fc = FormatConverter()
        results = fc.batch_convert_folder(str(source), str(output))
        assert results == []

    def test_batch_convert_with_xml_file(self, tmp_path):
        """Test batch conversion with an XML file."""
        source = tmp_path / "source"
        source.mkdir()
        xml_content = '<?xml version="1.0"?><root><data>test</data></root>'
        (source / "metadata.xml").write_text(xml_content, encoding="utf-8")
        output = tmp_path / "output"

        fc = FormatConverter()
        results = fc.batch_convert_folder(str(source), str(output))
        assert len(results) >= 1
        xml_results = [r for r in results if r.conversion_type == "xml_to_json"]
        assert len(xml_results) == 1
        assert xml_results[0].success is True

    def test_batch_convert_with_copy_file(self, tmp_path):
        """Test batch conversion with an unrecognized file type (copied as-is)."""
        source = tmp_path / "source"
        source.mkdir()
        (source / "readme.txt").write_text("hello", encoding="utf-8")
        output = tmp_path / "output"

        fc = FormatConverter()
        results = fc.batch_convert_folder(str(source), str(output))
        assert len(results) >= 1
        copy_results = [r for r in results if r.conversion_type == "copy"]
        assert len(copy_results) == 1

    def test_batch_convert_generates_manifest(self, tmp_path):
        """Test that batch conversion generates a manifest file."""
        source = tmp_path / "source"
        source.mkdir()
        (source / "test.txt").write_text("data", encoding="utf-8")
        output = tmp_path / "output"

        fc = FormatConverter()
        fc.batch_convert_folder(str(source), str(output))

        manifest_path = output / "conversion_manifest.json"
        assert manifest_path.exists()
        with open(manifest_path) as f:
            manifest = json.load(f)
        assert "conversion_date" in manifest
        assert manifest["total_files"] >= 1

    def test_batch_convert_with_fcs_file(self, tmp_path):
        """Test batch conversion with an FCS file."""
        source = tmp_path / "source"
        source.mkdir()

        keywords = {
            "$PAR": "1",
            "$TOT": "1",
            "$P1N": "FSC-A",
            "$BYTEORD": "1,2,3,4",
            "$DATATYPE": "I",
            "$MODE": "L",
            "$P1B": "16",
        }
        data = struct.pack("<H", 500)
        fcs_bytes = _build_minimal_fcs(keywords, data)
        (source / "sample.fcs").write_bytes(fcs_bytes)

        output = tmp_path / "output"
        fc = FormatConverter()
        results = fc.batch_convert_folder(str(source), str(output))
        fcs_results = [r for r in results if r.conversion_type == "fcs_to_csv"]
        assert len(fcs_results) == 1
        assert fcs_results[0].success is True

    def test_batch_convert_skips_lock_files(self, tmp_path):
        """Test that .lock files are skipped."""
        source = tmp_path / "source"
        source.mkdir()
        (source / "test.lock").write_text("lock", encoding="utf-8")
        output = tmp_path / "output"

        fc = FormatConverter()
        results = fc.batch_convert_folder(str(source), str(output))
        # Lock files should be skipped, so no results for them
        assert len(results) == 0


# ---------------------------------------------------------------------------
# _parse_fcs_data_segment tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestParseFcsDataSegment:
    """Tests for FormatConverter._parse_fcs_data_segment."""

    def test_double_data_type(self, tmp_path):
        """Test parsing FCS with $DATATYPE=D (64-bit double)."""
        keywords = {
            "$PAR": "2",
            "$TOT": "1",
            "$P1N": "FSC-A",
            "$P2N": "SSC-A",
            "$BYTEORD": "1,2,3,4",
            "$DATATYPE": "D",
            "$MODE": "L",
        }
        data = struct.pack("<dd", 1.5, 2.5)
        fcs_bytes = _build_minimal_fcs(keywords, data)

        fcs_path = tmp_path / "double.fcs"
        fcs_path.write_bytes(fcs_bytes)
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        csv_path, meta_path = fc.convert_fcs_to_csv(str(fcs_path), str(output_dir))
        assert csv_path != ""
        content = Path(csv_path).read_text()
        assert "FSC-A" in content

    def test_ascii_data_type(self, tmp_path):
        """Test parsing FCS with $DATATYPE=A (ASCII)."""
        keywords = {
            "$PAR": "2",
            "$TOT": "1",
            "$P1N": "FSC-A",
            "$P2N": "SSC-A",
            "$BYTEORD": "1,2,3,4",
            "$DATATYPE": "A",
            "$MODE": "L",
        }
        data = b"100 200"
        fcs_bytes = _build_minimal_fcs(keywords, data)

        fcs_path = tmp_path / "ascii.fcs"
        fcs_path.write_bytes(fcs_bytes)
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        csv_path, meta_path = fc.convert_fcs_to_csv(str(fcs_path), str(output_dir))
        assert csv_path != ""
        content = Path(csv_path).read_text()
        assert "FSC-A" in content

    def test_unknown_data_type(self, tmp_path):
        """Test parsing FCS with unknown $DATATYPE."""
        keywords = {
            "$PAR": "2",
            "$TOT": "1",
            "$P1N": "FSC-A",
            "$P2N": "SSC-A",
            "$BYTEORD": "1,2,3,4",
            "$DATATYPE": "Z",
            "$MODE": "L",
        }
        fcs_bytes = _build_minimal_fcs(keywords)

        fcs_path = tmp_path / "unknown_dtype.fcs"
        fcs_path.write_bytes(fcs_bytes)
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        csv_path, meta_path = fc.convert_fcs_to_csv(str(fcs_path), str(output_dir))
        assert csv_path != ""

    def test_non_list_mode(self, tmp_path):
        """Test parsing FCS with $MODE != L (unsupported)."""
        keywords = {
            "$PAR": "2",
            "$TOT": "1",
            "$P1N": "FSC-A",
            "$P2N": "SSC-A",
            "$BYTEORD": "1,2,3,4",
            "$DATATYPE": "I",
            "$MODE": "C",
            "$P1B": "16",
            "$P2B": "16",
        }
        fcs_bytes = _build_minimal_fcs(keywords)

        fcs_path = tmp_path / "nonlist.fcs"
        fcs_path.write_bytes(fcs_bytes)
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        csv_path, meta_path = fc.convert_fcs_to_csv(str(fcs_path), str(output_dir))
        assert csv_path != ""

    def test_supplemental_data_offsets(self, tmp_path):
        """Test FCS with header data offsets=0 but $BEGINDATA/$ENDDATA keywords."""
        keywords = {
            "$PAR": "1",
            "$TOT": "1",
            "$P1N": "FSC-A",
            "$BYTEORD": "1,2,3,4",
            "$DATATYPE": "I",
            "$MODE": "L",
            "$P1B": "16",
            "$BEGINDATA": "200",
            "$ENDDATA": "202",
        }
        # Build with data_start=0, data_end=0 in header
        text_segment = _build_fcs_text_segment(keywords)
        text_start = 58
        text_end = text_start + len(text_segment)
        # Data at offset 200
        data = struct.pack("<H", 500)
        padding = b"\x00" * (200 - text_end)
        header = _build_fcs_header(text_start, text_end, 0, 0)
        fcs_bytes = header + text_segment + padding + data

        fcs_path = tmp_path / "supplemental.fcs"
        fcs_path.write_bytes(fcs_bytes)
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        csv_path, meta_path = fc.convert_fcs_to_csv(str(fcs_path), str(output_dir))
        assert csv_path != ""


# ---------------------------------------------------------------------------
# _find_existing_tiff tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestFindExistingTiff:
    """Tests for FormatConverter._find_existing_tiff."""

    def test_finds_tif(self, tmp_path):
        czi = tmp_path / "image.czi"
        czi.write_bytes(b"dummy")
        tif = tmp_path / "image.tif"
        tif.write_bytes(b"tiff_data")

        fc = FormatConverter()
        result = fc._find_existing_tiff(czi)
        assert result == tif

    def test_finds_tiff(self, tmp_path):
        czi = tmp_path / "image.czi"
        czi.write_bytes(b"dummy")
        tiff = tmp_path / "image.tiff"
        tiff.write_bytes(b"tiff_data")

        fc = FormatConverter()
        result = fc._find_existing_tiff(czi)
        assert result == tiff

    def test_no_tiff_found(self, tmp_path):
        czi = tmp_path / "image.czi"
        czi.write_bytes(b"dummy")

        fc = FormatConverter()
        result = fc._find_existing_tiff(czi)
        assert result is None

    def test_empty_tiff_not_returned(self, tmp_path):
        czi = tmp_path / "image.czi"
        czi.write_bytes(b"dummy")
        tif = tmp_path / "image.tif"
        tif.write_bytes(b"")  # empty file

        fc = FormatConverter()
        result = fc._find_existing_tiff(czi)
        assert result is None


# ---------------------------------------------------------------------------
# convert_czi_to_tiff tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestConvertCziToTiff:
    """Tests for FormatConverter.convert_czi_to_tiff."""

    def test_convert_czi_placeholder(self, tmp_path):
        """Test CZI→TIFF creates placeholder when czifile not installed."""
        czi_path = tmp_path / "image.czi"
        czi_path.write_bytes(b"ZISRAW" + b"\x00" * 200)
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        results = fc.convert_czi_to_tiff(str(czi_path), str(output_dir))
        assert len(results) >= 1
        assert Path(results[0]).exists()

    def test_convert_czi_reuses_existing_tiff(self, tmp_path):
        """Test CZI→TIFF reuses existing TIFF alongside CZI."""
        czi_path = tmp_path / "image.czi"
        czi_path.write_bytes(b"ZISRAW" + b"\x00" * 200)
        tiff_path = tmp_path / "image.tif"
        tiff_path.write_bytes(b"II\x2a\x00" + b"\x00" * 100)
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        results = fc.convert_czi_to_tiff(str(czi_path), str(output_dir))
        assert len(results) >= 1
        # Should reuse existing TIFF
        assert str(tiff_path) in results


# ---------------------------------------------------------------------------
# convert_flowjo_to_json tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestConvertFlowjoToJson:
    """Tests for FormatConverter.convert_flowjo_to_json."""

    def test_convert_wsp_to_json(self, tmp_path):
        """Test WSP→JSON conversion."""
        wsp_content = (
            '<?xml version="1.0"?><Workspace><SampleList>'
            "<Sample>test</Sample></SampleList></Workspace>"
        )
        wsp_path = tmp_path / "test.wsp"
        wsp_path.write_text(wsp_content, encoding="utf-8")
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        json_path = fc.convert_flowjo_to_json(str(wsp_path), str(output_dir))
        assert json_path != ""
        assert Path(json_path).exists()

        with open(json_path) as f:
            data = json.load(f)
        assert "SampleList" in data

    def test_convert_wsp_nonexistent(self, tmp_path):
        """Test convert_flowjo_to_json with nonexistent file."""
        fc = FormatConverter()
        json_path = fc.convert_flowjo_to_json(str(tmp_path / "missing.wsp"), str(tmp_path / "out"))
        assert json_path == ""


# ---------------------------------------------------------------------------
# convert_ndpi_to_tiff tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestConvertNdpiToTiff:
    """Tests for FormatConverter.convert_ndpi_to_tiff."""

    def test_convert_ndpi_placeholder(self, tmp_path):
        """Test NDPI→TIFF creates placeholder when tifffile not installed."""
        ndpi_path = tmp_path / "slide.ndpi"
        ndpi_path.write_bytes(b"II\x2a\x00" + b"\x00" * 200)
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        tiff_path, meta_path = fc.convert_ndpi_to_tiff(str(ndpi_path), str(output_dir))
        assert tiff_path != ""
        assert Path(tiff_path).exists()
        assert meta_path != ""

    def test_convert_ndpi_nonexistent(self, tmp_path):
        """Test convert_ndpi_to_tiff with nonexistent file.

        Note: the method creates a placeholder TIFF even for missing files
        because tifffile/openslide both fail and the fallback writes a placeholder.
        """
        fc = FormatConverter()
        tiff_path, meta_path = fc.convert_ndpi_to_tiff(
            str(tmp_path / "missing.ndpi"), str(tmp_path / "out")
        )
        assert tiff_path != ""
        assert Path(tiff_path).exists()


# ---------------------------------------------------------------------------
# convert_lif_to_tiff tests
# ---------------------------------------------------------------------------


@pytest.mark.unit
class TestConvertLifToTiff:
    """Tests for FormatConverter.convert_lif_to_tiff."""

    def test_convert_lif_placeholder(self, tmp_path):
        """Test LIF→TIFF creates placeholder when readlif not installed."""
        lif_path = tmp_path / "confocal.lif"
        lif_path.write_bytes(b"\x00" * 200)
        output_dir = tmp_path / "output"

        fc = FormatConverter()
        tiff_path, meta_path = fc.convert_lif_to_tiff(str(lif_path), str(output_dir))
        assert tiff_path != ""
        assert Path(tiff_path).exists()

    def test_convert_lif_nonexistent(self, tmp_path):
        """Test convert_lif_to_tiff with nonexistent file.

        Note: the method creates a placeholder TIFF even for missing files.
        """
        fc = FormatConverter()
        tiff_path, meta_path = fc.convert_lif_to_tiff(
            str(tmp_path / "missing.lif"), str(tmp_path / "out")
        )
        assert tiff_path != ""
        assert Path(tiff_path).exists()
